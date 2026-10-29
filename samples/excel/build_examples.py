"""
Builds the example Excel workbooks in this folder. Invented data only.

    python build_examples.py                 # default sizes
    python build_examples.py --contracts 400 # a bigger clean batch for volume testing

Needs openpyxl (pip install openpyxl). Headers come from the canonical model, so the workbooks stay
in step with the upload template. Every cell is written as text, the way the template asks, so
leading zeros on CIFs and account numbers survive.
"""
import argparse
import csv
import json
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
MODEL = json.loads((ROOT / "ferry-server/src/main/resources/canonical-model.json").read_text(encoding="utf-8"))
SHEETS = [("CONTRACT", "contracts"), ("ACCOUNT", "accounts"), ("USER", "users"), ("SIGNATORY_GROUP", "signatory_groups")]

HEAD_FILL = PatternFill("solid", fgColor="161B26")
HEAD_FONT = Font(bold=True, color="FFFFFF")
REQ_FONT = Font(bold=True, color="F28500")


def fields(record_type):
    return MODEL["recordTypes"][record_type]["fields"]


def sheet_name(record_type):
    return MODEL["recordTypes"][record_type]["label"]


def add_sheet(wb, record_type, rows, header_override=None):
    """rows are dicts keyed by canonical field name (or by header_override's column names)."""
    ws = wb.create_sheet(sheet_name(record_type))
    if header_override:
        headers, keys = header_override, header_override
    else:
        headers = [f["label"] + (" *" if f.get("required") else "") for f in fields(record_type)]
        keys = [f["name"] for f in fields(record_type)]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill, c.font = HEAD_FILL, (REQ_FONT if h.endswith("*") else HEAD_FONT)
    for r, row in enumerate(rows, 2):
        for col, k in enumerate(keys, 1):
            v = row.get(k)
            if v not in (None, ""):
                c = ws.cell(row=r, column=col, value=str(v))
                c.number_format = "@"
    for col, h in enumerate(headers, 1):
        width = max([len(h)] + [len(str(row.get(keys[col - 1]) or "")) for row in rows])
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = min(max(width + 2, 10), 40)
    ws.freeze_panes = "A2"
    return ws


def add_readme(wb, lines):
    ws = wb.create_sheet("Readme", 0)   # sheets named Readme are ignored on upload
    ws.column_dimensions["A"].width = 120
    for i, line in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=line)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            c.font = Font(bold=True, size=14, color="F28500")


# ---------------------------------------------------------------- 1. the synthetic test batch

def synthetic(out):
    """The six contracts in samples/synthetic, as one workbook. Each one trips specific checks."""
    by_label = {}
    for rt, _ in SHEETS:
        by_label[rt] = {f["label"]: f["name"] for f in fields(rt)}
    wb = Workbook()
    wb.remove(wb.active)
    add_readme(wb, [
        "Synthetic test batch: every Boarding check in six contracts",
        "Upload the whole workbook in one go (What the file holds: Whole workbook). Leave the batch's source system blank.",
        "Expected plan: R-001 PARTIAL · R-002 READY · B-100 PARTIAL · B-200 PARTIAL · R-003 BLOCKED · R-004 BLOCKED.",
        "Why each contract ends that way is in samples/README.md. Invented data only.",
    ])
    for rt, stem in SHEETS:
        with open(ROOT / "samples/synthetic" / f"{stem}.csv", newline="", encoding="utf-8") as f:
            rows = [{by_label[rt][k]: v for k, v in row.items()} for row in csv.DictReader(f)]
        add_sheet(wb, rt, rows)
    wb.save(out)


# ---------------------------------------------------------------- 2. a clean batch at volume

FIRST = ["Alex", "Maya", "Noah", "Lena", "Omar", "Priya", "Sam", "Ines", "Kenji", "Zara", "Leo", "Nadia",
         "Ivan", "Chloe", "Ravi", "Elena", "Tariq", "Sofia", "Jonas", "Amara"]
MIDDLE = ["J", "Marie", "Lee", "Rose", "Ali", "Kai", "Ann", "Ray"]
LAST = ["Hart", "Okafor", "Silva", "Tanaka", "Novak", "Reyes", "Khan", "Moreau", "Berg", "Costa", "Adeyemi", "Lund"]
CITIES = [("GB", "London", "44"), ("SG", "Singapore", "65"), ("AE", "Dubai", "971"), ("US", "Austin", "1"), ("DE", "Berlin", "49")]
CURRENCIES = ["USD", "EUR", "GBP", "SGD"]
COMPANIES = ["Northwind", "Bluefin", "Harbor", "Summit", "Keystone", "Orchid", "Granite", "Lighthouse", "Meridian", "Cobalt"]


def clean(out, contracts, seed=7):
    """Retail and business contracts that should all plan as READY: the happy path, at any size."""
    rnd = random.Random(seed)
    c_rows, a_rows, u_rows, g_rows = [], [], [], []
    n_business = max(1, contracts // 4)
    seq = iter(range(1, 10_000_000))

    def cif():
        return f"0{next(seq):08d}"           # leading zero on purpose

    def account():
        return f"0{rnd.randrange(10**12, 10**13)}"

    def phone():
        return str(rnd.randrange(10**8, 10**9))

    def dob():
        return f"{rnd.randint(1958, 2002)}-{rnd.randint(1, 12):02d}-{rnd.randint(1, 28):02d}"

    def person():
        f, l = rnd.choice(FIRST), rnd.choice(LAST)
        return f, rnd.choice(MIDDLE), l

    for i in range(contracts):
        business = i < n_business
        country, city, cc = rnd.choice(CITIES)
        ref = f"B-{2001 + i}" if business else f"R-{1001 + i - n_business}"
        primary = cif()
        first, middle, last = person()
        company = f"{rnd.choice(COMPANIES)} {rnd.choice(['Trading', 'Logistics', 'Foods', 'Labs', 'Partners'])}"
        domain = (company.split()[0] if business else last).lower() + ".example.test"
        c_rows.append({
            "contractRef": ref, "serviceDefinitionRef": "BUSINESS" if business else "RETAIL", "primaryCif": primary,
            "email": f"{'finance' if business else first.lower()}@{domain}", "phoneCountryCode": cc, "phoneNumber": phone(),
            "addressLine1": f"{rnd.randint(1, 250)} {rnd.choice(['Harbour', 'Market', 'Station', 'Park', 'Mill'])} Road",
            "city": city, "country": country, "zipCode": str(rnd.randint(10000, 99999)),
        })
        cifs = [primary] + ([cif()] if business and rnd.random() < 0.4 else [])
        for c in cifs:
            for k in range(rnd.randint(1, 2) + (1 if business else 0)):
                a_rows.append({
                    "contractRef": ref, "cif": c, "accountId": account(),
                    "accountName": rnd.choice(["Operating", "Payroll", "Reserve"] if business else ["Savings", "Current", "Travel"]),
                    "accountType": rnd.choice(["C", "S"]), "currency": rnd.choice(CURRENCIES), "status": "Open",
                })

        def user(suffix, role, user_type, user_cif, who):
            f, m, l = who
            u_rows.append({
                "userRef": f"U-{ref}-{suffix}", "contractRef": ref, "roleRef": role, "userType": user_type, "cif": user_cif,
                "legacyUserName": f"{f.lower()}.{l.lower()}{rnd.randint(1, 99)}", "firstName": f, "middleName": m,
                "lastName": l, "email": f"{f.lower()}.{l.lower()}@{domain}", "phoneCountryCode": cc, "phone": phone(),
                "dateOfBirth": dob(), "taxId": f"TX{rnd.randrange(10**7, 10**8)}",
            })
            return f"U-{ref}-{suffix}"

        if not business:
            user("1", "USER", "LINKED", primary, (first, middle, last))
            continue
        admin = user("ADM", "ADMIN", "LINKED", primary, person())
        makers = [user(f"M{k}", "ACCOUNTANT", "VIRTUAL", "", person()) for k in range(1, rnd.randint(2, 3))]
        checkers = [user(f"C{k}", "ACCOUNTANT", "VIRTUAL", "", person()) for k in range(1, 3)]
        # One group per user on a contract (vendor rule); names letters, digits and spaces only.
        for member in [admin] + checkers:
            g_rows.append({"groupRef": f"G-{ref}-A", "contractRef": ref, "groupName": "Senior Approvers",
                           "description": "Approves payments above the threshold", "userRef": member})
        for member in makers:
            g_rows.append({"groupRef": f"G-{ref}-B", "contractRef": ref, "groupName": "Operations",
                           "description": "", "userRef": member})

    wb = Workbook()
    wb.remove(wb.active)
    add_readme(wb, [
        f"Clean batch: {contracts} contracts ({n_business} business, {contracts - n_business} retail)",
        "Upload the whole workbook in one go. Leave the batch's source system blank.",
        "Expected: every record valid, no errors, and every contract READY in the plan. Use it to see a clean run "
        "and how the screens behave with more data. Rebuild bigger with: python build_examples.py --contracts 400",
        "Business contracts have a linked admin, one or two virtual makers and two virtual checkers, in two signatory groups. "
        "Retail contracts have one linked user. All CIFs and account numbers start with 0. Invented data only.",
    ])
    for (rt, _), rows in zip(SHEETS, [c_rows, a_rows, u_rows, g_rows]):
        add_sheet(wb, rt, rows)
    wb.save(out)
    return len(c_rows), len(a_rows), len(u_rows), len(g_rows)


# ---------------------------------------------------------------- 3. a source system with its own columns

LEGACY_USER_COLUMNS = ["LOGIN_ID", "CONTRACT_NO", "ROLE_CODE", "USER_KIND", "CUST_NO", "FIRST_NM", "MIDDLE_NM", "LAST_NM",
                       "EMAIL_ADDR", "MOBILE", "BIRTH_DT", "NATIONAL_ID", "BRANCH_CODE"]


def legacycore(out):
    """Contracts and accounts in template headers; users exactly as the fictional LEGACYCORE exports them."""
    contracts = [
        {"contractRef": "L-501", "serviceDefinitionRef": "RETAIL", "primaryCif": "000510001", "email": "ida.park@example.test",
         "phoneCountryCode": "95", "phoneNumber": "912345670", "city": "Port Town", "country": "MM"},
        {"contractRef": "L-502", "serviceDefinitionRef": "RETAIL", "primaryCif": "000520002"},
        {"contractRef": "L-601", "serviceDefinitionRef": "BUSINESS", "primaryCif": "006010001", "email": "accounts@quay.example.test",
         "phoneCountryCode": "95", "phoneNumber": "912345671", "addressLine1": "4 Quay Lane", "city": "Port Town", "country": "MM"},
    ]
    accounts = [
        {"contractRef": "L-501", "cif": "000510001", "accountId": "05100000100001", "accountName": "Savings", "currency": "USD", "status": "Open"},
        {"contractRef": "L-502", "cif": "000520002", "accountId": "05200000200001", "accountName": "Current", "currency": "USD", "status": "Open"},
        {"contractRef": "L-601", "cif": "006010001", "accountId": "06010000100001", "accountName": "Operating", "currency": "USD", "status": "Open"},
        {"contractRef": "L-601", "cif": "006010001", "accountId": "06010000100002", "accountName": "Payroll", "currency": "EUR", "status": "Open"},
    ]
    # What the export looks like: C/N for linked/virtual, one-letter role codes, dd/MM/yyyy dates,
    # phones with dashes and spaces, mixed-case emails, and a BRANCH_CODE column nobody maps.
    users = [
        ["ida.park", "L-501", "U", "C", "000510001", "Ida", "", "Park", "Ida.Park@Example.test", "912-345-670", "14/02/1987", "", "051"],
        ["leo.grant", "L-502", "U", "C", "000520002", "Leo", "", "Grant", "LEO.GRANT@example.test", "(91) 234 5672", "30/11/1979", "", "052"],
        ["quay.admin", "L-601", "A", "C", "006010001", "Rosa", "", "Vane", "rosa@quay.example.test", "912 345 673", "01/06/1975", "", "060"],
        ["quay.maker", "L-601", "F", "N", "", "Tom", "Eli", "Shore", "tom@quay.example.test", "9123-45674", "22/09/1992", "NID-4471902", "060"],
        ["quay.checker", "L-601", "F", "N", "", "Asha", "Mae", "Reed", "asha@quay.example.test", "912 345 675", "05/03/1988", "NID-5520413", "060"],
    ]
    wb = Workbook()
    wb.remove(wb.active)
    add_readme(wb, [
        "LEGACYCORE batch: a source system with its own column names",
        "Create the batch with source system LEGACYCORE, then upload the whole workbook in one go.",
        "Contracts and Accounts use the template headers. Users is the raw LEGACYCORE export: the mapping profile "
        "config/mapping-profiles/example-legacycore-user.json renames its columns, turns C/N into LINKED/VIRTUAL, maps role "
        "codes A/U/F, keeps only the digits of phone numbers, lower-cases emails and converts dd/MM/yyyy dates.",
        "Expected: BRANCH_CODE listed as an ignored column, no errors, and all three contracts READY. Invented data only.",
    ])
    add_sheet(wb, "CONTRACT", contracts)
    add_sheet(wb, "ACCOUNT", accounts)
    add_sheet(wb, "USER", [dict(zip(LEGACY_USER_COLUMNS, u)) for u in users], header_override=LEGACY_USER_COLUMNS)
    wb.save(out)


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--contracts", type=int, default=40, help="contracts in clean-batch.xlsx (default 40)")
    args = ap.parse_args()
    synthetic(HERE / "synthetic-batch.xlsx")
    counts = clean(HERE / "clean-batch.xlsx", args.contracts)
    legacycore(HERE / "legacycore-batch.xlsx")
    print("synthetic-batch.xlsx   6 contracts (one per check)")
    print("clean-batch.xlsx       %d contracts, %d accounts, %d users, %d group members" % counts)
    print("legacycore-batch.xlsx  3 contracts, users in LEGACYCORE columns")
